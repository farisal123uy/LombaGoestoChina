import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "submission" / "generated"
DATA = json.loads((OUTPUT / "deck-data.json").read_text(encoding="utf-8"))
PPTX = OUTPUT / "Guardian-Wallet-presentation.pptx"

COLORS = {
    "paper": "F7F8F5",
    "white": "FFFFFF",
    "ink": "20312F",
    "muted": "5C706A",
    "green": "087F68",
    "green_dark": "125E4D",
    "green_pale": "E9F5EF",
    "green_line": "BB DCC C".replace(" ", ""),
    "violet": "70549F",
    "violet_pale": "F3EFF9",
    "amber": "865713",
    "amber_pale": "FFF5E0",
    "line": "DCE5DF",
    "card": "FBFCFA",
    "blue": "2C6587",
}


def rgb(name):
    return RGBColor.from_string(COLORS.get(name, name))


def add_shape(slide, x, y, w, h, fill, line=None, rounded=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(0.8)
    return shape


def add_text(slide, x, y, w, h, value, size, color="ink", bold=False, font="Aptos", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(value)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_link(slide, x, y, w, label, url):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.28))
    box.text_frame.margin_left = 0
    box.text_frame.margin_right = 0
    paragraph = box.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = label
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = rgb("green_dark")
    run.hyperlink.address = url


def add_bullet(slide, x, y, w, value, size=17):
    add_shape(slide, x, y + 0.12, 0.09, 0.09, "green", "green", rounded=True)
    add_text(slide, x + 0.25, y, w - 0.25, 0.57, value, size, "ink")


def add_header(slide, index, item):
    add_shape(slide, 0.58, 0.38, 0.42, 0.42, "green", "green")
    add_text(slide, 0.625, 0.41, 0.33, 0.31, "G", 18, "white", True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.12, 0.38, 2.5, 0.24, "GUARDIAN WALLET", 11, "ink", True)
    add_text(slide, 1.12, 0.64, 3.6, 0.18, "AI WITH BOUNDED AUTHORITY", 7.5, "muted", True)
    add_text(slide, 8.2, 0.43, 4.55, 0.25, item["eyebrow"], 9, "green", True, align=PP_ALIGN.RIGHT)
    add_shape(slide, 0.58, 0.97, 12.17, 0.012, "line", "line", rounded=False)


def add_aside(slide, item):
    x, y, w, h = 8.55, 2.55, 4.2, 3.74
    add_shape(slide, x, y, w, h, "green_pale", "green_line")
    add_text(slide, x + 0.3, y + 0.28, w - 0.6, 0.34, item["asideTitle"], 10, "green_dark", True)
    cursor = y + 0.85
    for bullet in item.get("aside", []):
        add_shape(slide, x + 0.33, cursor + 0.08, 0.07, 0.07, "green", "green")
        add_text(slide, x + 0.53, cursor, w - 0.83, 0.55, bullet, 13, "ink", True)
        cursor += 0.66
    links = item.get("links", [])
    if links:
        cursor = max(cursor + 0.08, y + h - 0.92)
        for link in links:
            add_link(slide, x + 0.32, cursor, w - 0.64, link["label"], link["url"])
            cursor += 0.3


def add_slide(presentation, index, item):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("paper")
    add_header(slide, index, item)

    add_text(slide, 0.7, 1.25, 7.55, 0.92, item["title"], 30, "ink", True)
    add_text(slide, 0.73, 2.12, 7.42, 0.55, item["subtitle"], 13, "muted")
    add_shape(slide, 0.7, 2.55, 7.55, 3.74, "white", "line")
    cursor = 2.88
    count = len(item["bullets"])
    bullet_size = 16 if count <= 4 else 13 if count >= 6 else 14
    step = 0.78 if count <= 4 else 0.53 if count >= 6 else 0.68
    for bullet in item["bullets"]:
        add_bullet(slide, 1.02, cursor, 6.95, bullet, bullet_size)
        cursor += step

    add_aside(slide, item)
    add_shape(slide, 0.7, 6.6, 12.05, 0.012, "line", "line", rounded=False)
    add_text(slide, 0.72, 6.78, 10.8, 0.28, item["evidence"], 9, "muted")
    add_text(slide, 11.3, 6.78, 1.35, 0.28, f"{index:02d} / {len(DATA['slides']):02d}", 9, "green_dark", True, align=PP_ALIGN.RIGHT)
    notes = f"{item['narration']}\n\nTiming target: {item['seconds']} seconds.\n\nEvidence links:\n" + "\n".join(link["url"] for link in item.get("links", []))
    if not item.get("links"):
        notes += f"\n{DATA['payment']['explorer']}"
    slide.notes_slide.notes_text_frame.text = notes
    return slide


OUTPUT.mkdir(parents=True, exist_ok=True)
presentation = Presentation()
presentation.core_properties.author = "Guardian Wallet team"
presentation.core_properties.title = DATA["title"]
presentation.core_properties.subject = f"{DATA['track']} · approximately {DATA['durationMinutes']} minutes"
presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
for index, item in enumerate(DATA["slides"], 1):
    add_slide(presentation, index, item)
presentation.save(PPTX)
print(json.dumps({"pptx": str(PPTX), "slides": len(presentation.slides), "targetMinutes": DATA["durationMinutes"]}, ensure_ascii=False))
