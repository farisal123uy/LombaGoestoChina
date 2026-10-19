import json
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "submission" / "generated"
data = json.loads((OUT / "deck-data.json").read_text(encoding="utf-8"))
presentation = Presentation(str(OUT / "Guardian-Wallet-presentation.pptx"))
expected_seconds = sum(int(item["seconds"]) for item in data["slides"])
assert len(presentation.slides) == len(data["slides"]), "Deck slide count differs from source data"
assert expected_seconds < 10 * 60, "Talk timing must fit under ten minutes"

all_text = []
note_status = []
hyperlinks = []
for slide, source in zip(presentation.slides, data["slides"], strict=True):
    slide_text = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
    all_text.append(slide_text)
    notes = slide.notes_slide.notes_text_frame.text
    assert source["narration"] in notes, f"Speaker notes missing for slide: {source['title']}"
    note_status.append(bool(notes.strip()))
    assert "Not a public" not in slide_text or "Local" in slide_text, "Fork evidence must be labeled locally"
    for shape in slide.shapes:
        if shape.left < 0 or shape.top < 0 or shape.left + shape.width > presentation.slide_width or shape.top + shape.height > presentation.slide_height:
            raise AssertionError(f"Shape exceeds slide canvas: {source['title']}")
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.hyperlink.address:
                        hyperlinks.append(run.hyperlink.address)

deck_text = "\n".join(all_text)
for required in ["Base Sepolia", "MockUSD", "60/60", "index 1", "Escrow release"]:
    assert required in deck_text, f"Required current fact missing from slides: {required}"
for key in ["deploymentExplorer", "lockExplorer", "releaseExplorer"]:
    assert data["payment"][key] in hyperlinks, f"Evidence link missing from slide hyperlinks: {key}"
assert all(note_status), "Every slide needs speaker notes"
notes_id = OUT / ".." / "presentation-notes.md"
assert notes_id.resolve().exists(), "Indonesian rehearsal notes missing"

print(json.dumps({
    "validated": True,
    "slides": len(presentation.slides),
    "timingSeconds": expected_seconds,
    "timingMinutes": round(expected_seconds / 60, 2),
    "speakerNotes": len(note_status),
    "clickableEvidenceLinks": len(hyperlinks),
    "pptxBytes": (OUT / "Guardian-Wallet-presentation.pptx").stat().st_size,
}, ensure_ascii=False))
